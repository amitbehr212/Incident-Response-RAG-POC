# Copyright 2025
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""File parsers for incident response documents."""

import io
from typing import Any

import docx
import fitz  # PyMuPDF
import openpyxl
from PIL import Image


def parse_pdf(file_content: bytes, file_name: str) -> str:
    """Parse PDF file using PyMuPDF.

    Args:
        file_content: PDF file content as bytes
        file_name: Name of the file for logging

    Returns:
        Extracted text from the PDF
    """
    try:
        doc = fitz.open(stream=file_content, filetype="pdf")
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        return text.strip()
    except Exception as e:
        print(f"Error parsing PDF {file_name}: {e}")
        return ""


def parse_word_document(file_content: bytes, file_name: str) -> str:
    """Parse Word document using python-docx.

    Args:
        file_content: Word file content as bytes
        file_name: Name of the file for logging

    Returns:
        Extracted text from the document
    """
    try:
        doc = docx.Document(io.BytesIO(file_content))
        text = "\n".join([para.text for para in doc.paragraphs])
        return text.strip()
    except Exception as e:
        print(f"Error parsing Word document {file_name}: {e}")
        return ""


def parse_excel(file_content: bytes, file_name: str) -> str:
    """Parse Excel file using openpyxl.

    Args:
        file_content: Excel file content as bytes
        file_name: Name of the file for logging

    Returns:
        Extracted text from all sheets
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        text = ""
        for sheet in wb.worksheets:
            text += f"\n=== Sheet: {sheet.title} ===\n"
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                if row_text.strip():
                    text += row_text + "\n"
        return text.strip()
    except Exception as e:
        print(f"Error parsing Excel {file_name}: {e}")
        return ""


def parse_plain_text(file_content: bytes, file_name: str) -> str:
    """Parse plain text file.

    Args:
        file_content: Text file content as bytes
        file_name: Name of the file for logging

    Returns:
        Decoded text content
    """
    try:
        return file_content.decode("utf-8", errors="ignore").strip()
    except Exception as e:
        print(f"Error parsing text file {file_name}: {e}")
        return ""


def parse_image(file_content: bytes, file_name: str) -> str:
    """Parse image file using OCR (pytesseract).

    Args:
        file_content: Image file content as bytes
        file_name: Name of the file for logging

    Returns:
        Extracted text from image via OCR
    """
    try:
        import pytesseract

        image = Image.open(io.BytesIO(file_content))
        text = pytesseract.image_to_string(image)
        return text.strip()
    except Exception as e:
        print(f"Error parsing image {file_name}: {e}")
        return ""


def parse_google_doc(drive_service: Any, file_id: str, file_name: str) -> str:
    """Parse Google Doc by exporting as plain text.

    Args:
        drive_service: Google Drive API service
        file_id: Google Drive file ID
        file_name: Name of the file for logging

    Returns:
        Extracted text from Google Doc
    """
    try:
        request = drive_service.files().export_media(fileId=file_id, mimeType="text/plain")
        file_content = request.execute()
        return parse_plain_text(file_content, file_name)
    except Exception as e:
        print(f"Error parsing Google Doc {file_name}: {e}")
        return ""


def parse_google_sheet(drive_service: Any, file_id: str, file_name: str) -> str:
    """Parse Google Sheet by exporting as Excel.

    Args:
        drive_service: Google Drive API service
        file_id: Google Drive file ID
        file_name: Name of the file for logging

    Returns:
        Extracted text from Google Sheet
    """
    try:
        request = drive_service.files().export_media(
            fileId=file_id, mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        file_content = request.execute()
        return parse_excel(file_content, file_name)
    except Exception as e:
        print(f"Error parsing Google Sheet {file_name}: {e}")
        return ""


def get_file_type_config() -> dict[str, dict[str, Any]]:
    """Get configuration mapping MIME types to parser functions.

    Returns:
        Dictionary mapping MIME types to parser configuration
    """
    return {
        # PDF files
        "application/pdf": {
            "parser": parse_pdf,
            "needs_drive": False,
            "display_name": "PDF",
        },
        # Word documents
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": {
            "parser": parse_word_document,
            "needs_drive": False,
            "display_name": "Word Document",
        },
        "application/msword": {
            "parser": parse_word_document,
            "needs_drive": False,
            "display_name": "Word Document (Legacy)",
        },
        # Excel spreadsheets
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {
            "parser": parse_excel,
            "needs_drive": False,
            "display_name": "Excel Spreadsheet",
        },
        "application/vnd.ms-excel": {
            "parser": parse_excel,
            "needs_drive": False,
            "display_name": "Excel Spreadsheet (Legacy)",
        },
        # Plain text
        "text/plain": {
            "parser": parse_plain_text,
            "needs_drive": False,
            "display_name": "Text File",
        },
        "text/markdown": {
            "parser": parse_plain_text,
            "needs_drive": False,
            "display_name": "Markdown File",
        },
        # Images
        "image/png": {
            "parser": parse_image,
            "needs_drive": False,
            "display_name": "PNG Image",
        },
        "image/jpeg": {
            "parser": parse_image,
            "needs_drive": False,
            "display_name": "JPEG Image",
        },
        # Google Workspace files
        "application/vnd.google-apps.document": {
            "parser": parse_google_doc,
            "needs_drive": True,
            "display_name": "Google Doc",
        },
        "application/vnd.google-apps.spreadsheet": {
            "parser": parse_google_sheet,
            "needs_drive": True,
            "display_name": "Google Sheet",
        },
    }
